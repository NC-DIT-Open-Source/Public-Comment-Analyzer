"""HTTP contract tests using real SQLite, files, authentication, and LangGraph.

Explicit demo inference verifies plumbing without claims about model quality.
The provider-adapter tests separately exercise the real LangChain boundary.
"""

from concurrent.futures import ThreadPoolExecutor
import csv
import io
import json
from pathlib import Path
import sys
import time
from urllib.parse import parse_qs, urlsplit
import uuid

import bcrypt
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
import pytest

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "backend" / "shared"))

from runtime import StorageError
from adapters.local import LocalRuntime


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("APP_RUNTIME", "local")
    monkeypatch.setenv("APP_DATA_DIR", str(tmp_path))
    monkeypatch.setenv("APP_ALLOWED_HOSTS", "testserver")
    monkeypatch.setenv("LLM_PROVIDER", "demo")
    monkeypatch.delenv("ALLOWED_ORIGIN", raising=False)
    monkeypatch.delenv("LLM_API_KEY_FILE", raising=False)
    monkeypatch.delenv("ACCESS_PASSWORD_HASH_FILE", raising=False)
    monkeypatch.setenv("LOCAL_PASSWORD_HASH", bcrypt.hashpw(b"contract-test", bcrypt.gensalt(rounds=4)).decode())
    monkeypatch.setenv("APP_MAX_ANALYSES_PER_MINUTE", "100")
    monkeypatch.setenv("LLM_MAX_CALLS", "1000")
    from backend.local.app import create_app
    app = create_app(start_worker=False)
    with TestClient(app) as http:
        http.headers["X-Access-Key"] = "contract-test"
        yield http


def _upload(client, rows=3, extension="csv"):
    if extension == "csv":
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(["id", "comment", "empty"])
        writer.writerows([[str(i), f"Public comment {i}, with café and a newline\nline two", ""] for i in range(rows)])
        content = output.getvalue().encode()
    else:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["id", "comment", "empty"])
        for i in range(rows):
            sheet.append([str(i), f"Public comment {i} café", ""])
        output = io.BytesIO()
        workbook.save(output)
        content = output.getvalue()
    response = client.post("/api/upload", files={"file": (f"comments.{extension}", content)})
    assert response.status_code == 200, response.text
    assert response.json()["rowCount"] == rows
    return response.json()


def test_demo_labels_use_the_same_normalized_provider_as_inference(client, monkeypatch):
    monkeypatch.setenv('LLM_PROVIDER', ' demo ')
    from backend.row_processor.handler import _demo_notice_column
    assert client.get('/api/config').json() == {'demoMode': True}
    assert _demo_notice_column(['comment'], []) == '_analysis_notice'


def _start(client, file_id, categorized=False):
    col = {"name": "Finding", "type": "open_text", "instructions": "Summarize the comment"}
    if categorized:
        col.update(type="categorized", options=[{"value": "Support", "description": "Support"}, {"value": "Oppose", "description": "Oppose"}])
    response = client.post("/api/process", json={
        "fileId": file_id, "selectedCommentColumn": "comment", "contextDescription": "Synthetic public comments",
        "analysisColumns": [col],
    })
    assert response.status_code == 200, response.text
    return response.json()["jobId"]


@pytest.mark.parametrize("extension", ["csv", "xlsx"])
def test_upload_process_summary_download_roundtrip(client, extension):
    uploaded = _upload(client, extension=extension)
    job_id = _start(client, uploaded["fileId"])
    runtime = client.app.state.runtime
    assert client.get(f"/api/status/{job_id}").json()["status"] == "pending"
    assert runtime.run_one()
    status = client.get(f"/api/status/{job_id}").json()
    assert status["status"] == "completed"
    assert status["totalRows"] == status["completedRows"] == 3
    pending_summary = client.get(f"/api/results/{job_id}").json()
    assert pending_summary["aggregateAnalysis"] is None
    assert pending_summary["analysisStatus"] == "generating"
    # Navigation downloads have no shared-password header.
    result = client.get(pending_summary["downloadUrl"], headers={"X-Access-Key": ""})
    assert result.status_code == 200
    assert "attachment" in result.headers["content-disposition"]
    if extension == "csv":
        rows = list(csv.DictReader(io.StringIO(result.text)))
        assert list(rows[0]) == ["id", "comment", "empty", "Finding", "_error", "_analysis_notice"]
        assert 'category values are placeholders' in rows[0]['_analysis_notice']
        assert [r["id"] for r in rows] == ["0", "1", "2"]
        assert rows[0]["comment"].endswith("\nline two")
        assert rows[0]["empty"] == rows[0]["_error"] == ""
    else:
        workbook = load_workbook(io.BytesIO(result.content), read_only=True)
        rows = list(workbook.active.values)
        assert rows[0] == ("id", "comment", "empty", "Finding", "_error", "_analysis_notice")
        assert 'no AI inference' in rows[1][-1]
        assert len(rows) == 4
        workbook.close()
    assert runtime.run_one()
    summary = client.get(f"/api/results/{job_id}").json()
    assert "Demo mode" in summary["aggregateAnalysis"]
    dashboard = client.post(f"/api/dashboard/{job_id}", json={"prompt": "Show a chart"})
    assert dashboard.status_code == 200
    assert dashboard.json()["charts"] == []
    assert "Demo mode" in dashboard.json()["narrative"]


def test_categorized_preview_confirmation_preserves_contract(client):
    file_id = _upload(client, rows=50)["fileId"]
    job_id = _start(client, file_id, categorized=True)
    runtime = client.app.state.runtime
    assert runtime.run_one()
    preview = client.get(f"/api/status/{job_id}").json()
    assert preview["status"] == "preview_ready"
    assert len(preview["previewRows"]) == 20
    assert preview["totalRows"] == 50
    assert preview["selectedCommentColumn"] == "comment"
    assert preview["analysisColumns"][0]["options"][0]["value"] == "Support"
    assert 'category values are placeholders' in preview['previewRows'][0]['_analysis_notice']
    assert not runtime.run_one()  # Paid work must wait for the person.
    assert client.post(f"/api/process/{job_id}/preview-confirm", json={}).status_code == 200
    assert client.post(f"/api/process/{job_id}/preview-confirm", json={}).status_code == 409
    assert runtime.run_one()
    complete = client.get(f"/api/status/{job_id}").json()
    assert complete["status"] == "completed" and complete["completedRows"] == 50


def test_http_rejects_internal_async_flags_and_auth_bypass(client):
    assert client.post("/api/process", json={"asyncProcessing": True}).status_code == 400
    for path in ("/api/upload", "/api/process", f"/api/dashboard/{uuid.uuid4()}"):
        response = client.post(path, json={"asyncProcessing": True}, headers={"X-Access-Key": ""})
        assert response.status_code == 401
    assert client.post("/api/process/extra/path", json={}).status_code == 404
    assert client.get("/api/config").json() == {"demoMode": True}


def test_download_signature_scope_tamper_expiry_and_private_uploads(client):
    objects = client.app.state.runtime.objects
    job_id = str(uuid.uuid4())
    key = f"results/{job_id}/output.csv"
    objects.put(key, b"result\nverified")
    signed = objects.signed_url(key)
    query = parse_qs(urlsplit(signed).query)
    assert client.get(signed, headers={"X-Access-Key": ""}).status_code == 200
    params = {k: values[0] for k, values in query.items()}
    params["signature"] = "0" * 64
    assert client.get("/api/download", params=params).status_code == 403
    params['signature'] = 'é'
    assert client.get('/api/download', params=params).status_code == 403
    params["expires"] = int(time.time()) - 1
    assert client.get("/api/download", params=params).status_code == 403
    with pytest.raises(StorageError):
        objects.signed_url(f"uploads/{uuid.uuid4()}/input.csv")
    for key in ("../../etc/passwd", f"results/{job_id}/../../secret", f"results/{job_id}/output.csv/extra"):
        with pytest.raises(StorageError):
            objects.put(key, b"rejected")


def test_jobs_survive_restart_and_interrupted_paid_work_is_not_replayed(client):
    file_id = _upload(client)["fileId"]
    job_id = _start(client, file_id)
    old = client.app.state.runtime
    task = old._claim()
    assert task
    restarted = LocalRuntime(old.root)
    restarted.recover()
    assert restarted.jobs.get(job_id)["status"] == "failed"
    assert not restarted.run_one()


def test_runtime_lifecycle_stops_new_inference_and_resumes_after_start(tmp_path, monkeypatch):
    from inference import _enabled, InferenceLimitError, pause_for_shutdown, resume_after_startup
    monkeypatch.setenv('LLM_ENABLED', 'true')
    monkeypatch.delenv('LLM_KILL_SWITCH_FILE', raising=False)
    runtime = LocalRuntime(tmp_path)
    pause_for_shutdown()
    try:
        runtime.start()
        _enabled()
        runtime.stop()
        with pytest.raises(InferenceLimitError, match='shuts down'):
            _enabled()
        assert not runtime.healthy()
    finally:
        runtime.stop()
        resume_after_startup()


@pytest.mark.parametrize('phase', ['full', 'confirm', 'preview'])
def test_recovery_preserves_saved_results_before_queue_acknowledgment(client, phase):
    runtime = client.app.state.runtime
    job_id = str(uuid.uuid4())
    key = f'results/{job_id}/output.csv'
    item = {'jobId': job_id, 'status': 'completed', 'outputFileKey': key}
    if phase == 'preview':
        item.update(status='preview_ready', previewRows=[{'comment': 'synthetic', 'Position': 'Support'}])
    else:
        runtime.objects.put(key, b'comment,Position\nsynthetic,Support')
    runtime.jobs.put(item)
    runtime.enqueue('row_processor', {'jobId': job_id, 'phase': phase})
    assert runtime._claim()
    runtime.recover()
    assert runtime.jobs.get(job_id)['status'] == item['status']
    assert not runtime.run_one()
    if phase != 'preview':
        result = client.get(f'/api/results/{job_id}')
        assert result.status_code == 200
        assert client.get(result.json()['downloadUrl']).status_code == 200


def test_recovery_preserves_committed_summary_without_replaying(client):
    runtime = client.app.state.runtime
    job_id = str(uuid.uuid4())
    runtime.jobs.put({'jobId': job_id, 'status': 'completed', 'aggregateAnalysis': 'Saved summary'})
    runtime.enqueue('aggregate_analyzer', {'pathParameters': {'jobId': job_id}})
    assert runtime._claim()
    runtime.recover()
    assert runtime.jobs.get(job_id)['aggregateAnalysis'] == 'Saved summary'
    assert runtime.jobs.get(job_id).get('analysisStatus') != 'failed'
    assert not runtime.run_one()


def test_preview_claim_is_atomic(client):
    jobs = client.app.state.runtime.jobs
    job_id = str(uuid.uuid4())
    jobs.put({"jobId": job_id, "status": "preview_ready"})
    with ThreadPoolExecutor(max_workers=8) as pool:
        results = list(pool.map(lambda _: jobs.claim_preview(job_id), range(20)))
    assert results.count(True) == 1


def test_cancel_pending_job_releases_slot_and_never_runs(client):
    job_id = _start(client, _upload(client)['fileId'])
    assert client.post(f'/api/process/{job_id}/cancel', json={}, headers={'X-Access-Key': ''}).status_code == 401
    response = client.post(f'/api/process/{job_id}/cancel', json={})
    assert response.status_code == 200
    assert client.get(f'/api/status/{job_id}').json()['status'] == 'failed'
    assert not client.app.state.runtime.run_one()


def test_recovery_fails_orphan_jobs_and_expires_abandoned_previews(client):
    runtime = client.app.state.runtime
    orphan = str(uuid.uuid4())
    expired = str(uuid.uuid4())
    runtime.jobs.put({'jobId': orphan, 'status': 'processing'})
    runtime.jobs.put({'jobId': expired, 'status': 'preview_ready', 'ttl': int(time.time()) - 1})
    runtime.recover()
    assert runtime.jobs.get(orphan)['status'] == 'failed'
    assert runtime.jobs.get(expired)['status'] == 'failed'


def test_malformed_shapes_body_caps_and_headers(client):
    assert client.post('/api/process', json=[]).status_code == 400
    assert client.post('/api/auth/validate', content=b'x' * 1025).status_code == 413
    result = client.get('/api/health')
    assert result.headers['cache-control'] == 'no-store'
    assert "script-src 'self'" in result.headers['content-security-policy']
    assert result.headers['x-content-type-options'] == 'nosniff'


@pytest.mark.parametrize('name,selected', [('comment', 'comment'), ('COMMENT', 'comment'), ('Finding', 'typo')])
def test_invalid_column_selection_rejected_before_job_or_model_work(client, name, selected):
    file_id = _upload(client)['fileId']
    response = client.post('/api/process', json={
        'fileId': file_id, 'selectedCommentColumn': selected, 'contextDescription': 'Synthetic comments',
        'analysisColumns': [{'name': name, 'type': 'open_text', 'instructions': 'Summarize'}],
    })
    assert response.status_code == 400
    assert response.json()['error']['code'] == 'INVALID_COLUMNS'
    assert not client.app.state.runtime.run_one()


def test_source_error_column_is_preserved_by_rejecting_ambiguous_export(client):
    response = client.post('/api/upload', files={'file': ('comments.csv', b'comment,_error\nHello,original note')})
    assert response.status_code == 200
    file_id = response.json()['fileId']
    result = client.post('/api/process', json={
        'fileId': file_id, 'selectedCommentColumn': 'comment', 'contextDescription': 'Synthetic comments',
        'analysisColumns': [{'name': 'Finding', 'type': 'open_text', 'instructions': 'Summarize'}],
    })
    assert result.status_code == 400
    assert client.app.state.runtime.objects.path(f'uploads/{file_id}/input.csv').read_bytes().endswith(b'original note')


@pytest.mark.parametrize('extension', ['csv', 'xlsx'])
@pytest.mark.parametrize('source_columns', [3, 1000])
def test_formula_header_and_wide_source_keep_summary_dashboard_counts(client, monkeypatch, extension, source_columns):
    headers = ['comment'] + [f'metadata{i}' for i in range(source_columns - 1)]
    rows = [[f'Synthetic comment {i}'] + ['original'] * (source_columns - 1) for i in range(2)]
    if extension == 'csv':
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(headers)
        writer.writerows(rows)
        content = output.getvalue().encode()
    else:
        workbook = Workbook()
        workbook.active.append(headers)
        for row in rows:
            workbook.active.append(row)
        output = io.BytesIO()
        workbook.save(output)
        content = output.getvalue()
    upload = client.post('/api/upload', files={'file': (f'comments.{extension}', content)})
    assert upload.status_code == 200
    response = client.post('/api/process', json={
        'fileId': upload.json()['fileId'], 'selectedCommentColumn': 'comment', 'contextDescription': 'Synthetic comments',
        'analysisColumns': [{'name': '+Finding', 'type': 'categorized', 'instructions': 'Classify',
                             'options': [{'value': 'Support', 'description': 'Support'}, {'value': 'Oppose', 'description': 'Oppose'}]}],
    })
    assert response.status_code == 200, response.text
    job_id = response.json()['jobId']
    runtime = client.app.state.runtime
    assert runtime.run_one()
    assert runtime.jobs.get(job_id)['status'] == 'completed'
    assert runtime.jobs.get(job_id)['exportHeaders'][source_columns] == '+Finding'
    from backend.aggregate_analyzer import handler as aggregate
    from backend.dashboard_generator import handler as dashboard
    prompts = {}
    def summarize(prompt):
        prompts['summary'] = prompt
        return 'Draft summary for synthetic test'
    def charts(prompt):
        prompts['dashboard'] = prompt
        return json.dumps({'charts': [], 'narrative': 'Demo dashboard'})
    monkeypatch.setattr(aggregate, '_call_summary_model', summarize)
    monkeypatch.setattr(dashboard, '_call_dashboard_model', charts)
    assert runtime.run_one()
    result = client.get(f'/api/results/{job_id}').json()
    assert result['aggregateAnalysis'] == 'Draft summary for synthetic test'
    assert 'Support: 2 (100.0%)' in prompts['summary']
    assert client.post(f'/api/dashboard/{job_id}', json={'prompt': 'Count positions'}).status_code == 200
    assert 'Support: 2 (100.0%)' in prompts['dashboard']
    assert 'Synthetic comment 0' in prompts['dashboard']
    download = client.get(result['downloadUrl'])
    assert download.status_code == 200
    if extension == 'csv':
        exported = list(csv.reader(io.StringIO(download.text)))
    else:
        workbook = load_workbook(io.BytesIO(download.content), read_only=True, data_only=False)
        exported = list(workbook.active.values)
        workbook.close()
    assert exported[0][source_columns] == "'+Finding"
    assert exported[1][source_columns] == 'Support'
    assert exported[1][0] == 'Synthetic comment 0'


@pytest.mark.parametrize('source_headers,analysis_name', [
    (['comment', "'+Finding"], '+Finding'),
    (['comment', '+metadata', "'+metadata"], 'Finding'),
])
def test_escaped_header_collision_is_rejected_before_job_and_inference(client, monkeypatch, source_headers, analysis_name):
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(source_headers)
    writer.writerow(['synthetic'] * len(source_headers))
    upload = client.post('/api/upload', files={'file': ('comments.csv', output.getvalue().encode())})
    assert upload.status_code == 200
    from backend.row_processor import handler
    monkeypatch.setattr(handler, 'preflight_job', lambda *_args, **_kwargs: pytest.fail('Inference preflight must not run for ambiguous headers'))
    response = client.post('/api/process', json={
        'fileId': upload.json()['fileId'], 'selectedCommentColumn': 'comment', 'contextDescription': 'Synthetic comments',
        'analysisColumns': [{'name': analysis_name, 'type': 'open_text', 'instructions': 'Summarize'}],
    })
    assert response.status_code == 400
    assert response.json()['error']['code'] == 'INVALID_COLUMNS'
    assert 'collide' in response.json()['error']['message']
    assert not client.app.state.runtime.run_one()
    with client.app.state.runtime.jobs.connect() as connection:
        assert connection.execute('SELECT COUNT(*) FROM jobs').fetchone()[0] == 0


@pytest.mark.parametrize('extension', ['csv', 'xlsx'])
def test_formula_prefixed_categories_keep_exact_summary_and_dashboard_counts(client, monkeypatch, extension):
    uploaded = _upload(client, rows=2, extension=extension)
    columns = [{'name': '+Position', 'type': 'categorized', 'instructions': 'Classify',
                'options': [{'value': '-1', 'description': 'Negative'}, {'value': '+1', 'description': 'Positive'}]}]
    response = client.post('/api/process', json={
        'fileId': uploaded['fileId'], 'selectedCommentColumn': 'comment',
        'contextDescription': 'Synthetic comments', 'analysisColumns': columns,
    })
    assert response.status_code == 200, response.text
    job_id = response.json()['jobId']
    from backend.row_processor import handler as rows
    from backend.aggregate_analyzer import handler as aggregate
    from backend.dashboard_generator import handler as dashboard
    monkeypatch.setattr(rows, '_process_single_row', lambda row, *_args: {'+Position': '-1' if row['id'] == '0' else '+1'})
    prompts = {}
    def summarize(prompt):
        prompts['summary'] = prompt
        return 'Synthetic draft summary'
    def charts(prompt):
        prompts['dashboard'] = prompt
        return json.dumps({'charts': [], 'narrative': 'Synthetic demo chart'})
    monkeypatch.setattr(aggregate, '_call_summary_model', summarize)
    monkeypatch.setattr(dashboard, '_call_dashboard_model', charts)
    runtime = client.app.state.runtime
    original_upload = runtime.objects.path(f"uploads/{uploaded['fileId']}/input.{extension}").read_bytes()
    assert runtime.run_one()
    assert runtime.run_one()
    result = client.get(f'/api/results/{job_id}').json()
    assert result['aggregateAnalysis'] == 'Synthetic draft summary'
    assert client.post(f'/api/dashboard/{job_id}', json={'prompt': 'Count each category'}).status_code == 200
    for prompt in prompts.values():
        assert '-1: 1 (50.0%)' in prompt
        assert '+1: 1 (50.0%)' in prompt
        assert '(unmatched/blank)' not in prompt
    download = client.get(result['downloadUrl'])
    if extension == 'csv':
        exported = list(csv.DictReader(io.StringIO(download.text)))
    else:
        workbook = load_workbook(io.BytesIO(download.content), read_only=True, data_only=False)
        assert all(cell.data_type != 'f' for row in workbook.active for cell in row)
        values = list(workbook.active.values)
        exported = [dict(zip(values[0], row)) for row in values[1:]]
        workbook.close()
    assert [row["'+Position"] for row in exported] == ["'-1", "'+1"]
    assert [row['id'] for row in exported] == ['0', '1']
    assert runtime.objects.path(f"uploads/{uploaded['fileId']}/input.{extension}").read_bytes() == original_upload


def test_escaped_category_collision_rejected_before_job_or_inference(client, monkeypatch):
    uploaded = _upload(client)
    from backend.row_processor import handler
    monkeypatch.setattr(handler, 'preflight_job', lambda *_args, **_kwargs: pytest.fail('Ambiguous categories must fail before inference preflight'))
    response = client.post('/api/process', json={
        'fileId': uploaded['fileId'], 'selectedCommentColumn': 'comment', 'contextDescription': 'Synthetic comments',
        'analysisColumns': [{'name': 'Position', 'type': 'categorized', 'instructions': 'Classify',
                             'options': [{'value': '-1', 'description': 'First'}, {'value': "'-1", 'description': 'Second'}]}],
    })
    assert response.status_code == 400
    assert response.json()['error']['code'] == 'INVALID_COLUMNS'
    assert 'Category labels collide' in response.json()['error']['message']
    assert not client.app.state.runtime.run_one()
    with client.app.state.runtime.jobs.connect() as connection:
        assert connection.execute('SELECT COUNT(*) FROM jobs').fetchone()[0] == 0


def test_demo_notice_does_not_overwrite_existing_source_column(client):
    response = client.post('/api/upload', files={'file': ('comments.csv', b'comment,_analysis_notice\nHello,original notice')})
    assert response.status_code == 200
    job_id = _start(client, response.json()['fileId'])
    assert client.app.state.runtime.run_one()
    result = client.get(f'/api/results/{job_id}').json()
    rows = list(csv.DictReader(io.StringIO(client.get(result['downloadUrl']).text)))
    assert rows[0]['_analysis_notice'] == 'original notice'
    assert 'no AI inference' in rows[0]['_analysis_notice_2']


def test_partial_row_failure_cannot_publish_completion_before_output_upload(client, monkeypatch):
    job_id = _start(client, _upload(client)['fileId'])
    runtime = client.app.state.runtime
    from backend.row_processor import handler
    original_classify = handler._process_single_row
    def classify(row, *args, **kwargs):
        if row['id'] == '1':
            raise RuntimeError('Synthetic row failure')
        return original_classify(row, *args, **kwargs)
    monkeypatch.setattr(handler, '_process_single_row', classify)
    original_upload = runtime.objects.upload
    original_update = runtime.jobs.update
    def upload(path, key):
        assert runtime.jobs.get(job_id)['status'] == 'processing'
        assert not runtime.objects.exists(key)
        return original_upload(path, key)
    def update(identity, fields):
        if fields.get('status') == 'completed':
            assert runtime.objects.exists(runtime.jobs.get(identity)['outputFileKey'])
        return original_update(identity, fields)
    monkeypatch.setattr(runtime.objects, 'upload', upload)
    monkeypatch.setattr(runtime.jobs, 'update', update)
    assert runtime.run_one()
    state = client.get(f'/api/status/{job_id}').json()
    assert state['status'] == 'completed'
    assert len(state['errors']) == 1
    result = client.get(f'/api/results/{job_id}').json()
    rows = list(csv.DictReader(io.StringIO(client.get(result['downloadUrl']).text)))
    assert len(rows) == 3 and rows[1]['_error'] and rows[0]['_error'] == ''


def test_failed_preview_rows_keep_preview_phase_until_saved(client, monkeypatch):
    job_id = _start(client, _upload(client, rows=50)['fileId'], categorized=True)
    runtime = client.app.state.runtime
    from backend.row_processor import handler
    monkeypatch.setattr(handler, '_process_single_row', lambda *_args, **_kwargs: (_ for _ in ()).throw(RuntimeError('Synthetic failure')))
    original_update = runtime.jobs.update
    states = []
    def update(identity, fields):
        if fields.get('status'):
            states.append(fields['status'])
        if fields.get('completedRows') and not fields.get('status'):
            assert runtime.jobs.get(identity)['status'] == 'preview_processing'
        return original_update(identity, fields)
    monkeypatch.setattr(runtime.jobs, 'update', update)
    assert runtime.run_one()
    assert states == ['preview_processing', 'preview_ready']
    assert len(runtime.jobs.get(job_id)['previewRows']) == 20
    assert not runtime.objects.exists(runtime.jobs.get(job_id)['outputFileKey'])


def test_auth_rate_limit_is_bounded_before_password_work(client):
    for _ in range(12):
        assert client.post('/api/auth/validate', json={'password': 'incorrect'}).status_code == 401
    assert client.post('/api/auth/validate', json={'password': 'incorrect'}).status_code == 429


def test_static_manifest_serves_build_and_spa_without_exposing_private_files(client, tmp_path, monkeypatch):
    from backend.local.app import create_app
    static = tmp_path / 'build'
    static.mkdir()
    (static / 'index.html').write_text('<html>Application shell</html>')
    (static / 'main.js').write_text('const application = true;')
    (static / 'assets').mkdir()
    (static / 'assets' / 'logo.svg').write_text('<svg></svg>')
    (static / '.env').write_text('synthetic private configuration')
    (static / '.private').mkdir()
    (static / '.private' / 'hidden.txt').write_text('synthetic hidden data')
    private = tmp_path / 'private'
    private.mkdir()
    (private / 'secret.txt').write_text('synthetic outside secret')
    (static / 'outside.txt').symlink_to(private / 'secret.txt')
    (static / 'linked-directory').symlink_to(private, target_is_directory=True)
    monkeypatch.setenv('APP_STATIC_DIR', str(static))
    with TestClient(create_app(start_worker=False)) as http:
        for path in ('/', '/upload', '/process/example'):
            response = http.get(path)
            assert response.status_code == 200
            assert response.text == '<html>Application shell</html>'
            assert response.headers['cache-control'] == 'no-cache'
        assert http.get('/index.html').headers['cache-control'] == 'no-cache'
        assert http.get('/main.js').text == 'const application = true;'
        assert http.get('/assets/logo.svg').text == '<svg></svg>'
        for path in ('/.env', '/.private/hidden.txt', '/outside.txt', '/linked-directory/secret.txt',
                     '/%2e%2e/private/secret.txt', '/assets/%2e%2e/.env', '/%2e%2e%5cprivate%5csecret.txt',
                     '/missing.js', '/%00'):
            response = http.get(path)
            assert response.status_code == 404, path
            assert 'synthetic' not in response.text
        # Build changes require restart; a request cannot discover extra files.
        (static / 'later.txt').write_text('not in the startup manifest')
        assert http.get('/later.txt').status_code == 404


def test_missing_static_build_keeps_local_api_available(client, tmp_path, monkeypatch):
    from backend.local.app import create_app
    monkeypatch.setenv('APP_STATIC_DIR', str(tmp_path / 'not-built-yet'))
    with TestClient(create_app(start_worker=False)) as http:
        assert http.get('/').status_code == 404
        assert http.get('/api/health').status_code == 200


def test_summary_failure_keeps_row_download_available(client, monkeypatch):
    job_id = _start(client, _upload(client)["fileId"])
    runtime = client.app.state.runtime
    assert runtime.run_one()
    from backend.aggregate_analyzer import handler
    monkeypatch.setattr(handler, "_call_summary_model", lambda *_: (_ for _ in ()).throw(RuntimeError("provider unavailable")))
    assert runtime.run_one()
    response = client.get(f"/api/results/{job_id}")
    assert response.status_code == 200
    assert response.json()["analysisStatus"] == "failed"
    assert client.get(response.json()["downloadUrl"]).status_code == 200


def test_auth_secret_file_rotation_and_startup_fail_closed(tmp_path, monkeypatch):
    monkeypatch.setenv("APP_RUNTIME", "local")
    monkeypatch.setenv("APP_DATA_DIR", str(tmp_path))
    monkeypatch.setenv("APP_ALLOWED_HOSTS", "testserver")
    monkeypatch.delenv("ALLOWED_ORIGIN", raising=False)
    monkeypatch.setenv("LLM_PROVIDER", "demo")
    monkeypatch.delenv("LOCAL_PASSWORD_HASH", raising=False)
    monkeypatch.delenv("ACCESS_PASSWORD_SECRET_NAME", raising=False)
    secret = tmp_path / "password.hash"
    monkeypatch.setenv("ACCESS_PASSWORD_HASH_FILE", str(secret))
    from backend.local.app import create_app
    from runtime import RuntimeConfigurationError
    with pytest.raises(RuntimeConfigurationError):
        with TestClient(create_app(start_worker=False)):
            pass
    secret.write_text(bcrypt.hashpw(b"first", bcrypt.gensalt(rounds=4)).decode())
    with TestClient(create_app(start_worker=False)) as http:
        assert http.post("/api/auth/validate", json={"password": "first"}).status_code == 200
        secret.write_text(bcrypt.hashpw(b"second", bcrypt.gensalt(rounds=4)).decode())
        assert http.post("/api/auth/validate", json={"password": "first"}).status_code == 401
        assert http.post("/api/auth/validate", json={"password": "second"}).status_code == 200
